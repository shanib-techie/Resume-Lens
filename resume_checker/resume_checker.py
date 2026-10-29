"""Resume text extraction + skill matching."""
import difflib
import io
import logging
import re
import unicodedata
import zipfile
from functools import lru_cache

from docx import Document
from pypdf import PdfReader

from .skills_db import ALIASES, ALL_SKILLS

log = logging.getLogger("resumelens")

GROUP_OF = {s.lower(): g for s, g in ALL_SKILLS}
SUPPORTED = (".pdf", ".docx")

MB = 1024 * 1024
MAX_RESUME_BYTES = 10 * MB        # one resume
MAX_UNZIPPED_BYTES = 60 * MB      # everything inside one ZIP
MAX_ZIP_ENTRIES = 100
MAX_PDF_PAGES = 30

# Words that are also common English words or names. They only match when written
# capitalised in the resume ("Go", "R", "Excel"), not as ordinary lowercase words.
AMBIGUOUS = {"go", "r", "c", "swift", "rust", "ruby", "excel", "express", "spring",
             "react", "sas", "dart", "notion", "blender", "maya", "sketch"}

FUZZY_MIN_LEN = 7
FUZZY_RATIO = 0.86


# ---------- extraction ----------
def extract_text(filename, data):
    name = filename.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                if not reader.decrypt(""):
                    raise ValueError("PDF is password protected")
            except ValueError:
                raise
            except Exception:
                raise ValueError("PDF is password protected")
        parts = []
        for page in reader.pages[:MAX_PDF_PAGES]:
            try:
                parts.append(page.extract_text() or "")
            except Exception:
                continue                           # skip an unreadable page, keep the rest
        return "\n".join(parts)
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.extend(c.text for c in row.cells)
        return "\n".join(parts)
    if name.endswith(".doc"):
        raise ValueError("Old .doc files are not supported. Save it as .docx or PDF")
    raise ValueError("Only PDF and DOCX files are supported")


# ---------- matching ----------
def normalize(text):
    text = unicodedata.normalize("NFKC", text).lower()   # fixes ligatures like "ﬁ"
    text = re.sub(r"[^a-z0-9+#.\s]", " ", text)           # keep c++, c#, .net, node.js
    return re.sub(r"\s+", " ", text).strip()


_ALIAS_N = {normalize(k): {normalize(a) for a in v} for k, v in ALIASES.items()}
_ALIAS_OF = {}                                            # alias -> {canonical names}
for _k, _alts in _ALIAS_N.items():
    for _a in _alts:
        _ALIAS_OF.setdefault(_a, set()).add(_k)


@lru_cache(maxsize=4096)
def _variants(skill):
    s = normalize(skill)
    out = {s, s.replace(" ", ""), s.replace(".", "")}
    out |= _ALIAS_N.get(s, set())
    for key in _ALIAS_OF.get(s, ()):                      # "powerbi" -> "power bi" and its siblings
        out.add(key)
        out.add(key.replace(" ", ""))
        out |= _ALIAS_N[key]
    return frozenset(v for v in out if v)


@lru_cache(maxsize=8192)
def _pattern(variant, case_sensitive):
    if case_sensitive:
        alts = "|".join(re.escape(a) for a in sorted({variant.capitalize(), variant.upper()}))
        return re.compile(rf"(?<![A-Za-z0-9+#.])(?:{alts})(?![A-Za-z0-9+#&-])")
    return re.compile(rf"(?<![a-z0-9+#]){re.escape(variant)}s?(?![a-z0-9+#])")


class _Doc:
    """One resume, prepared once and reused for every skill."""

    def __init__(self, text):
        self.raw = unicodedata.normalize("NFKC", text)
        self.norm = normalize(text)
        self.tokens = self.norm.split()
        self._grams = {}

    def grams(self, n):
        if n not in self._grams:
            t = self.tokens
            self._grams[n] = {" ".join(t[i:i + n]) for i in range(len(t) - n + 1)}
        return self._grams[n]


def _fuzzy(s, doc):
    n = len(s.split())
    for g in doc.grams(n):
        if abs(len(g) - len(s)) > 2 or g[0] != s[0]:
            continue
        sm = difflib.SequenceMatcher(None, s, g)
        if (sm.real_quick_ratio() >= FUZZY_RATIO and sm.quick_ratio() >= FUZZY_RATIO
                and sm.ratio() >= FUZZY_RATIO):
            return True
    return False


def match_skill(skill, doc):
    """Return 'exact', 'fuzzy' or None."""
    for v in _variants(skill):
        if v in AMBIGUOUS:
            if _pattern(v, True).search(doc.raw):
                return "exact"
        elif _pattern(v, False).search(doc.norm):
            return "exact"
    s = normalize(skill)
    if len(s) >= FUZZY_MIN_LEN and s not in AMBIGUOUS and _fuzzy(s, doc):   # typo tolerance
        return "fuzzy"
    return None


def analyze_text(text, skills):
    doc = _Doc(text)
    seen, matched, missing = set(), [], []
    for raw in skills:
        skill = raw.strip()
        if not skill or skill.lower() in seen:
            continue
        seen.add(skill.lower())
        hit = match_skill(skill, doc)
        item = {"skill": skill, "group": GROUP_OF.get(skill.lower(), "Custom")}
        if hit:
            item["match"] = hit
            matched.append(item)
        else:
            missing.append(item)
    total = len(matched) + len(missing)
    pct = round(len(matched) / total * 100, 1) if total else 0.0
    return {
        "total": total, "matched": matched, "missing": missing,
        "match_percent": pct, "verdict": verdict(pct),
        "suggestions": suggestions(missing, pct), "word_count": len(doc.tokens),
    }


def verdict(pct):
    if pct >= 80: return "Strong match"
    if pct >= 60: return "Good match"
    if pct >= 40: return "Partial match"
    return "Weak match"


def suggestions(missing, pct):
    tips, by_group = [], {}
    for m in missing:
        by_group.setdefault(m["group"], []).append(m["skill"])
    for g, items in by_group.items():
        tips.append(f"Add {', '.join(items[:6])} ({g}) to your skills or projects if you have real experience with them.")
    if pct < 60:
        tips.append("Use the exact wording from the job description so ATS filters can find it.")
    if not missing:
        tips.append("Every skill was found. Back them up with measurable results in your experience section.")
    return tips


# ---------- files ----------
def iter_resumes(filename, data):
    """Yield (name, bytes, error). A .zip upload is expanded into its resumes.
    Raises ValueError if the ZIP itself is unusable or too big."""
    if not filename.lower().endswith(".zip"):
        if len(data) > MAX_RESUME_BYTES:
            yield filename, None, f"Resume is larger than {MAX_RESUME_BYTES // MB} MB"
        else:
            yield filename, data, None
        return
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise ValueError("Not a valid ZIP file")
    with z:
        count = total = 0
        for info in z.infolist():
            path = info.filename.replace("\\", "/")
            base = path.rsplit("/", 1)[-1]
            if (info.is_dir() or not base or base.startswith((".", "~$")) or "__MACOSX" in path
                    or not base.lower().endswith(SUPPORTED)):
                continue
            count += 1
            total += info.file_size
            if count > MAX_ZIP_ENTRIES:
                raise ValueError(f"ZIP has more than {MAX_ZIP_ENTRIES} resumes")
            if total > MAX_UNZIPPED_BYTES:
                raise ValueError(f"ZIP expands to more than {MAX_UNZIPPED_BYTES // MB} MB")
            if info.flag_bits & 0x1:
                raise ValueError("ZIP is password protected")
            if info.file_size > MAX_RESUME_BYTES:
                yield base, None, f"Resume is larger than {MAX_RESUME_BYTES // MB} MB"
                continue
            yield base, z.read(info), None


def analyze_file(filename, data, skills):
    result = {"filename": filename, "file_type": filename.rsplit(".", 1)[-1].upper()}
    try:
        text = extract_text(filename, data)
        if not text.strip():
            raise ValueError("No readable text found (scanned PDF?)")
        result.update(analyze_text(text, skills))
    except ValueError as e:                          # our own, readable messages
        result["error"] = str(e)
    except Exception:                                # one bad file must not stop the batch
        log.exception("Could not read %s", filename)
        result["error"] = "Could not read this file (corrupt or unsupported)"
    return result
